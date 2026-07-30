from __future__ import annotations

from contextlib import contextmanager
from datetime import datetime
from pathlib import Path
from typing import Iterator


def _display_xls_value(book, cell):
    import xlrd

    if cell.ctype == xlrd.XL_CELL_DATE:
        value = xlrd.xldate_as_datetime(cell.value, book.datemode)
        return value.date() if value.time() == datetime.min.time() else value
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR}:
        return None
    return cell.value


@contextmanager
def read_spreadsheet_rows(path: str) -> Iterator[list[tuple[str, list[list[object]]]]]:
    """Read modern and legacy Excel workbooks into sheet names and row values."""
    extension = Path(path).suffix.lower()
    if extension == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise RuntimeError("Legacy .xls support requires the xlrd package") from exc

        book = xlrd.open_workbook(path, on_demand=True)
        try:
            sheets = []
            for sheet in book.sheets():
                rows = [
                    [_display_xls_value(book, sheet.cell(row_index, column_index)) for column_index in range(sheet.ncols)]
                    for row_index in range(sheet.nrows)
                ]
                sheets.append((sheet.name, rows))
            yield sheets
        finally:
            book.release_resources()
        return

    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        yield [
            (worksheet.title, [list(row) for row in worksheet.iter_rows(values_only=True)])
            for worksheet in workbook.worksheets
        ]
    finally:
        workbook.close()
