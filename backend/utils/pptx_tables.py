from __future__ import annotations

from math import ceil

from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


def add_paginated_table_slides(
    presentation,
    title: str,
    rows: list[list[str]],
    max_rows: int = 14,
    max_columns: int = 8,
) -> int:
    if not rows or not any(any(str(value).strip() for value in row) for row in rows):
        return 0

    column_count = max(len(row) for row in rows)
    normalized_rows = [list(map(str, row)) + [""] * (column_count - len(row)) for row in rows]
    header = normalized_rows[0]
    data_rows = normalized_rows[1:]

    slide_width = presentation.slide_width / 914400
    slide_height = presentation.slide_height / 914400
    table_left = 0.45
    table_top = 0.95
    table_width = max(1.0, slide_width - 0.9)
    table_height = max(1.0, slide_height - 1.35)
    if column_count >= 10:
        max_columns = max(max_columns, 10)
        max_rows = min(max_rows, 12)
    elif column_count >= 7:
        max_columns = max(max_columns, 8)
        max_rows = min(max_rows, 13)

    row_window = max(2, max_rows - 1)
    row_chunks = [data_rows[index:index + row_window] for index in range(0, len(data_rows), row_window)] or [[]]
    column_chunks = [list(range(index, min(index + max_columns, column_count))) for index in range(0, column_count, max_columns)]

    slide_count = 0
    total_pages = len(row_chunks) * len(column_chunks)
    for row_chunk in row_chunks:
        for columns in column_chunks:
            slide_count += 1
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            title_box = slide.shapes.add_textbox(Inches(table_left), Inches(0.22), Inches(table_width), Inches(0.55))
            title_paragraph = title_box.text_frame.paragraphs[0]
            title_paragraph.text = title if total_pages == 1 else f"{title} ({slide_count}/{total_pages})"
            title_paragraph.font.size = Pt(18)
            title_paragraph.font.bold = False

            page_rows = [header] + row_chunk
            table_shape = slide.shapes.add_table(
                len(page_rows), len(columns), Inches(table_left), Inches(table_top), Inches(table_width), Inches(table_height)
            )
            table = table_shape.table
            font_size = 10 if len(columns) <= 6 else (8 if len(columns) <= 10 else 6)
            for page_column in range(len(columns)):
                table.columns[page_column].width = int(table_shape.width / len(columns))
            for row_index, row in enumerate(page_rows):
                table.rows[row_index].height = int(table_shape.height / max(len(page_rows), 1))
                for page_column, source_column in enumerate(columns):
                    cell = table.cell(row_index, page_column)
                    cell.text = row[source_column]
                    cell.text_frame.word_wrap = True
                    cell.margin_left = Inches(0.04)
                    cell.margin_right = Inches(0.04)
                    cell.margin_top = Inches(0.03)
                    cell.margin_bottom = Inches(0.03)
                    for paragraph in cell.text_frame.paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(font_size)
                            run.font.bold = row_index == 0
                            if row_index == 0:
                                run.font.color.rgb = RGBColor(31, 41, 55)
    return slide_count
