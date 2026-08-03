from docx import Document
from docx.shared import Inches, Pt
from bs4 import BeautifulSoup, NavigableString, Tag
from io import BytesIO
import os
from .base import BaseConverter
from typing import Dict, Any
from backend.utils.html_resources import resolve_html_image
from backend.utils.html_content import prepare_content_soup
from backend.utils.text_utils import read_text_file, sanitize_xml_text
from .html_snapshot import render_html_snapshot, should_include_rendered_snapshot


class HtmlToDocxConverter(BaseConverter):
    """HTML 到 DOCX 转换器（优化版 - 参考 conversion_core）
    
    优化内容：
    1. 添加进度回调
    2. 支持两种模式：源码模式 / 文本提取模式
    3. 更好的文本清理
    """
    
    def __init__(self):
        super().__init__()
        self.supported_formats = ['docx', 'doc']

    @staticmethod
    def _clean_text(value: object, strip: bool = True) -> str:
        text = sanitize_xml_text(value)
        return text.strip() if strip else text

    def _add_paragraph(self, doc: Document, value: object, style: str | None = None):
        text = self._clean_text(value)
        if not text:
            return None
        if style:
            return doc.add_paragraph(text, style=style)
        return doc.add_paragraph(text)

    @staticmethod
    def _table_rows(table_element: Tag) -> list[list[str]]:
        rows = []
        for row in table_element.find_all('tr'):
            if row.find_parent('table') is not table_element:
                continue
            cells = [cell.get_text(' ', strip=True) for cell in row.find_all(['th', 'td'], recursive=False)]
            if cells:
                rows.append(cells)
        return rows

    def _render_table(self, doc: Document, element: Tag) -> None:
        rows = self._table_rows(element)
        column_count = max((len(row) for row in rows), default=0)
        if not rows or not column_count:
            return
        table = doc.add_table(rows=len(rows), cols=column_count)
        table.style = 'Table Grid'
        for row_index, row in enumerate(rows):
            for column_index, value in enumerate(row):
                table.cell(row_index, column_index).text = self._clean_text(value)

    def _render_image(self, doc: Document, element: Tag, base_path: str) -> bool:
        src = element.get('src')
        resolved = resolve_html_image(src, base_path=base_path)
        if not resolved:
            alt_text = str(element.get('alt') or '').strip()
            if alt_text:
                self._add_paragraph(doc, alt_text)
            return False

        image_bytes, _media_type = resolved
        if self._add_image_bytes(doc, image_bytes):
            return True

        alt_text = str(element.get('alt') or '').strip()
        if alt_text:
            self._add_paragraph(doc, alt_text)
        return False

    @staticmethod
    def _prepare_image_for_docx(image_bytes: bytes) -> tuple[bytes, float]:
        render_bytes = image_bytes
        width_inches = 5.8
        try:
            from PIL import Image

            with Image.open(BytesIO(image_bytes)) as image:
                width_inches = min(max(image.width / 96, 0.3), 5.8)
                if image.format not in {"PNG", "JPEG", "BMP", "GIF", "TIFF"}:
                    converted = BytesIO()
                    image.convert("RGBA" if "A" in image.getbands() else "RGB").save(converted, format="PNG")
                    render_bytes = converted.getvalue()
        except Exception:
            pass
        return render_bytes, width_inches

    def _add_image_bytes(self, doc: Document, image_bytes: bytes) -> bool:
        if not image_bytes:
            return False

        render_bytes, width_inches = self._prepare_image_for_docx(image_bytes)
        paragraph = doc.add_paragraph()
        run = paragraph.add_run()
        try:
            run.add_picture(BytesIO(render_bytes), width=Inches(width_inches))
            return True
        except Exception:
            paragraph._element.getparent().remove(paragraph._element)
            return False

    def _add_rendered_snapshot(self, doc: Document, input_path: str, output_path: str, options: dict) -> bool:
        try:
            output_dir = os.path.dirname(output_path) or os.getcwd()
            image_bytes = render_html_snapshot(input_path, output_dir, options)
            return self._add_image_bytes(doc, image_bytes)
        except Exception:
            return False

    def _render_list(self, doc: Document, element: Tag, base_path: str) -> None:
        style = 'List Bullet' if element.name == 'ul' else 'List Number'
        for item in element.find_all('li', recursive=False):
            text_parts = []
            nested_lists = []
            images = []
            for child in item.children:
                if isinstance(child, NavigableString):
                    text_parts.append(str(child).strip())
                elif isinstance(child, Tag) and child.name in {'ul', 'ol'}:
                    nested_lists.append(child)
                elif isinstance(child, Tag) and child.name == 'img':
                    images.append(child)
                elif isinstance(child, Tag):
                    text_parts.append(child.get_text(' ', strip=True))
            text = ' '.join(part for part in text_parts if part).strip()
            if text:
                self._add_paragraph(doc, text, style=style)
            for image in images:
                self._render_image(doc, image, base_path)
            for nested_list in nested_lists:
                self._render_list(doc, nested_list, base_path)

    def _render_text_and_images(self, doc: Document, element: Tag, base_path: str) -> None:
        text = element.get_text(' ', strip=True)
        if text:
            self._add_paragraph(doc, text)
        for image in element.find_all('img'):
            self._render_image(doc, image, base_path)

    def _render_blocks(self, doc: Document, parent: Tag, base_path: str) -> None:
        for element in parent.children:
            if isinstance(element, NavigableString):
                text = str(element).strip()
                if text:
                    self._add_paragraph(doc, text)
                continue
            if not isinstance(element, Tag):
                continue

            if element.name in {'h1', 'h2', 'h3', 'h4', 'h5', 'h6'}:
                text = self._clean_text(element.get_text(' ', strip=True))
                if text:
                    doc.add_heading(text, level=int(element.name[1]))
            elif element.name == 'table':
                self._render_table(doc, element)
            elif element.name in {'ul', 'ol'}:
                self._render_list(doc, element, base_path)
            elif element.name == 'img':
                self._render_image(doc, element, base_path)
            elif element.name in {'p', 'pre', 'blockquote'}:
                self._render_text_and_images(doc, element, base_path)
            elif element.find(['h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'p', 'pre', 'blockquote', 'table', 'ul', 'ol', 'img']):
                self._render_blocks(doc, element, base_path)
            else:
                text = element.get_text(' ', strip=True)
                if text:
                    self._add_paragraph(doc, text)
    
    def convert(self, input_path: str, output_path: str, **options) -> Dict[str, Any]:
        """将 HTML 转换为 DOCX"""
        try:
            self.validate_input(input_path)
            self.update_progress(input_path, 10)
            
            # 获取选项
            mode = options.get('mode', 'text')  # 'source' or 'text'
            
            # 读取 HTML
            html_content = read_text_file(input_path, options.get('encoding'))
            
            self.update_progress(input_path, 30)
            
            # 创建 Word 文档
            doc = Document()
            
            if mode == 'source':
                # 源码模式：保留完整 HTML 源码
                paragraph = doc.add_paragraph()
                run = paragraph.add_run(self._clean_text(html_content, strip=False))
                run.font.name = 'Consolas'
                run.font.size = Pt(10)
            else:
                include_snapshot = should_include_rendered_snapshot(
                    BeautifulSoup(html_content, "html.parser"),
                    options,
                )
                soup = prepare_content_soup(html_content, options)
                self.update_progress(input_path, 50)
                root = soup.body or soup
                if include_snapshot:
                    self._add_rendered_snapshot(doc, input_path, output_path, options)
                self._render_blocks(doc, root, input_path)
                self.update_progress(input_path, 80)
            
            # 保存文档
            doc.save(output_path)
            self.update_progress(input_path, 100)
            
            return {
                'success': True,
                'output_path': output_path,
                'size': self.get_output_size(output_path),
                'mode': mode
            }
            
        except Exception as e:
            self.cleanup_on_error(output_path)
            raise Exception(f"HTML to DOCX conversion failed: {str(e)}")
