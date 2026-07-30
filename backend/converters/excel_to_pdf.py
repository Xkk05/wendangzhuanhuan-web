from typing import Dict, Any
import os
import shutil
import logging
import tempfile
from pathlib import Path
from .base import BaseConverter
from backend.utils.logger import setup_logger

class ExcelToPdfConverter(BaseConverter):
    """Excel 转 PDF 转换器 - 优化版"""
    
    def __init__(self):
        super().__init__()
        self.supported_formats = ['pdf']
        self.soffice_path = self._find_libreoffice()
        self.logger = setup_logger('ExcelToPdf')

    def _find_libreoffice(self) -> str:
        """查找 LibreOffice 路径"""
        # 1. 优先尝试集成路径 (resources/libreoffice)
        from conversion_core.tools.office_to_pdf import get_bundled_libreoffice_path
        bundled_path = get_bundled_libreoffice_path()
        if bundled_path and os.path.exists(bundled_path):
            return bundled_path

        # 2. 尝试系统默认路径
        paths = [
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        ]
        for path in paths:
            if os.path.exists(path):
                return path
        
        # 3. 检查环境变量 PATH
        return shutil.which("soffice")

    def _resolve_page_settings(self, options: Dict[str, Any], col_count: int) -> Dict[str, Any]:
        """Return COM-ready page settings, preserving the existing automatic defaults."""
        paper_sizes = {
            'A4': 9,
            'A3': 8,
            'Letter': 1,
            'Legal': 5,
        }
        requested_orientation = str(options.get('excel_orientation') or 'auto').lower()
        scale_mode = str(options.get('excel_scale_mode') or 'auto').lower()

        if col_count <= 6:
            automatic = {'paper_size': 9, 'orientation': 1, 'scale_mode': 'fit_width'}
        elif col_count <= 10:
            automatic = {'paper_size': 9, 'orientation': 2, 'scale_mode': 'fit_width'}
        elif col_count <= 15:
            automatic = {'paper_size': 8, 'orientation': 2, 'scale_mode': 'fit_width'}
        else:
            automatic = {'paper_size': 8, 'orientation': 2, 'scale_mode': 'scaled_85'}

        if requested_orientation == 'portrait':
            automatic['orientation'] = 1
        elif requested_orientation == 'landscape':
            automatic['orientation'] = 2

        page_size = str(options.get('page_size') or 'auto')
        if page_size in paper_sizes:
            automatic['paper_size'] = paper_sizes[page_size]

        if scale_mode in {'fit_width', 'actual_size'}:
            automatic['scale_mode'] = scale_mode
        return automatic

    def _convert_with_excel(self, input_path: str, output_path: str, **options) -> Dict[str, Any]:
        """使用 Microsoft Excel 转换 - 优化版"""
        try:
            import comtypes.client
        except ImportError:
            raise Exception("comtypes库未安装，无法使用Excel COM转换")
        
        excel = None
        workbook = None
        
        try:
            # 转换为绝对路径
            input_path = os.path.abspath(input_path)
            output_path = os.path.abspath(output_path)
            
            self.logger.info(f"使用 Excel COM 转换: {input_path} -> {output_path}")
            
            # 启动 Excel
            excel = comtypes.client.CreateObject('Excel.Application')
            excel.Visible = False
            excel.DisplayAlerts = False
            try: excel.AutomationSecurity = 3  # msoAutomationSecurityForceDisable
            except: pass
            
            # 打开工作簿
            workbook = excel.Workbooks.Open(input_path, ReadOnly=True)
            
            # 优化每个工作表的页面设置
            for sheet in workbook.Worksheets:
                try:
                    # 1. 精准获取实际使用区域
                    used_range = None
                    try:
                        # 强制刷新 UsedRange
                        _ = sheet.UsedRange.Rows.Count
                        
                        # 精准查找最后一行和最后一列
                        last_row_cell = sheet.Cells.Find("*", SearchOrder=1, SearchDirection=2)
                        last_col_cell = sheet.Cells.Find("*", SearchOrder=2, SearchDirection=2)
                        
                        if last_row_cell and last_col_cell:
                            last_row = last_row_cell.Row
                            last_col = last_col_cell.Column
                            used_range = sheet.Range(sheet.Cells(1, 1), sheet.Cells(last_row, last_col))
                        else:
                            used_range = sheet.UsedRange
                    except:
                        used_range = sheet.UsedRange

                    if not used_range:
                        continue

                    # 2. 先自动调整列宽和行高，确保内容完整显示
                    try:
                        used_range.Columns.AutoFit()
                        used_range.Rows.AutoFit()
                        # 增加15%的宽度余量，防止文字被截断
                        for col_idx in range(1, min(used_range.Columns.Count + 1, 100)):
                            try:
                                col = used_range.Columns[col_idx]
                                curr_width = col.ColumnWidth
                                if 0 < curr_width < 100:
                                    col.ColumnWidth = curr_width * 1.15
                            except:
                                continue
                    except Exception as e:
                        print(f"列宽调整失败: {e}")

                    # 3. 使用用户页面设置；未指定时延续按列数的自动策略
                    setup = sheet.PageSetup
                    col_count = used_range.Columns.Count
                    page_settings = self._resolve_page_settings(options, col_count)
                    setup.PaperSize = page_settings['paper_size']
                    setup.Orientation = page_settings['orientation']

                    # 4. 设置缩放策略
                    if page_settings['scale_mode'] == 'fit_width':
                        setup.Zoom = False
                        setup.FitToPagesWide = 1
                        setup.FitToPagesTall = False
                    elif page_settings['scale_mode'] == 'actual_size':
                        setup.Zoom = 100
                        setup.FitToPagesWide = False
                        setup.FitToPagesTall = False
                    else:
                        setup.Zoom = 85
                        setup.FitToPagesWide = False
                        setup.FitToPagesTall = False
                    
                    # 5. 优化页边距（单位：磅）
                    try:
                        setup.LeftMargin = excel.Application.InchesToPoints(0.2)    # 约5mm
                        setup.RightMargin = excel.Application.InchesToPoints(0.2)
                        setup.TopMargin = excel.Application.InchesToPoints(0.3)     # 约7.5mm
                        setup.BottomMargin = excel.Application.InchesToPoints(0.3)
                        setup.HeaderMargin = excel.Application.InchesToPoints(0.1)
                        setup.FooterMargin = excel.Application.InchesToPoints(0.1)
                    except:
                        # 降级方案：直接使用磅值
                        setup.LeftMargin = 14
                        setup.RightMargin = 14
                        setup.TopMargin = 22
                        setup.BottomMargin = 22
                        setup.HeaderMargin = 7
                        setup.FooterMargin = 7
                    
                    # 6. 居中显示
                    setup.CenterHorizontally = True
                    setup.CenterVertically = False

                    # 7. 设置打印区域
                    try:
                        setup.PrintArea = used_range.Address
                    except:
                        pass
                    
                    # 8. 打印质量优化
                    setup.PrintGridlines = False  # 不打印网格线
                    setup.BlackAndWhite = False   # 彩色打印
                    try:
                        setup.PrintQuality = 600  # 高质量打印
                    except:
                        pass
                                
                except Exception as e_sheet:
                    print(f"工作表 {sheet.Name} 优化失败: {e_sheet}")
                    continue
            
            # 导出 PDF，使用最高质量设置
            try:
                workbook.ExportAsFixedFormat(
                    Type=0,  # xlTypePDF
                    Filename=output_path,
                    Quality=0,  # xlQualityStandard (最高质量)
                    IncludeDocProperties=True,
                    IgnorePrintAreas=False,
                    OpenAfterPublish=False
                )
            except Exception as e_export:
                print(f"ExportAsFixedFormat 失败: {e_export}")
                # 降级方案
                workbook.ExportAsFixedFormat(0, output_path)
            
            return {'method': 'microsoft_excel'}
            
        except Exception as e:
            self.logger.error(f"Excel COM 转换失败: {str(e)}")
            raise e
        finally:
            # 清理资源
            try:
                if workbook:
                    workbook.Close(SaveChanges=False)
            except:
                pass
            finally:
                workbook = None
            
            try:
                if excel:
                    excel.Quit()
            except:
                pass
            finally:
                excel = None

            # Release this conversion's COM references without terminating other
            # Excel instances that may contain unsaved user work.
            try:
                import gc
                gc.collect()
            except Exception:
                pass

    def _run_libreoffice_convert(
        self,
        input_path: str,
        target_format: str,
        output_dir: str,
        profile_dir: str,
    ) -> str:
        """Run one isolated LibreOffice conversion and return the generated path."""
        import subprocess

        os.makedirs(output_dir, exist_ok=True)
        os.makedirs(profile_dir, exist_ok=True)
        cmd = [
            self.soffice_path,
            '--headless',
            '--nologo',
            '--nodefault',
            '--nofirststartwizard',
            f'-env:UserInstallation={Path(profile_dir).resolve().as_uri()}',
            '--convert-to', target_format,
            '--outdir', output_dir,
            input_path
        ]

        result = subprocess.run(cmd, capture_output=True, text=True, timeout=60)
        if result.returncode != 0:
            detail = result.stderr.strip() or result.stdout.strip()
            raise Exception(f"LibreOffice conversion failed: {detail}")

        base_name = os.path.splitext(os.path.basename(input_path))[0]
        generated_path = os.path.join(output_dir, f"{base_name}.{target_format}")
        if not os.path.exists(generated_path):
            raise Exception(f"LibreOffice did not create the expected {target_format} file")
        return generated_path

    def _prepare_libreoffice_workbook(
        self,
        input_path: str,
        output_path: str,
        options: Dict[str, Any],
    ) -> None:
        """Write page settings into a temporary OOXML workbook for LibreOffice."""
        from openpyxl import load_workbook
        from openpyxl.worksheet.properties import PageSetupProperties

        keep_vba = os.path.splitext(input_path)[1].lower() == '.xlsm'
        workbook = load_workbook(input_path, keep_vba=keep_vba)
        try:
            for sheet in workbook.worksheets:
                page_settings = self._resolve_page_settings(options, max(sheet.max_column or 1, 1))
                page_setup = sheet.page_setup
                page_setup.paperSize = str(page_settings['paper_size'])
                page_setup.orientation = (
                    sheet.ORIENTATION_LANDSCAPE
                    if page_settings['orientation'] == 2
                    else sheet.ORIENTATION_PORTRAIT
                )

                setup_properties = sheet.sheet_properties.pageSetUpPr
                if setup_properties is None:
                    setup_properties = PageSetupProperties()
                    sheet.sheet_properties.pageSetUpPr = setup_properties

                if page_settings['scale_mode'] == 'fit_width':
                    setup_properties.fitToPage = True
                    page_setup.fitToWidth = 1
                    page_setup.fitToHeight = 0
                    page_setup.scale = None
                elif page_settings['scale_mode'] == 'actual_size':
                    setup_properties.fitToPage = False
                    page_setup.fitToWidth = None
                    page_setup.fitToHeight = None
                    page_setup.scale = 100
                else:
                    setup_properties.fitToPage = False
                    page_setup.fitToWidth = None
                    page_setup.fitToHeight = None
                    page_setup.scale = 85

                sheet.page_margins.left = 0.2
                sheet.page_margins.right = 0.2
                sheet.page_margins.top = 0.3
                sheet.page_margins.bottom = 0.3
                sheet.page_margins.header = 0.1
                sheet.page_margins.footer = 0.1
                sheet.print_options.horizontalCentered = True

            workbook.save(output_path)
        finally:
            workbook.close()

    def _convert_with_libreoffice(self, input_path: str, output_path: str, **options) -> Dict[str, Any]:
        """Use LibreOffice after applying the requested workbook page settings."""
        if not self.soffice_path:
            raise Exception("LibreOffice not found")

        output_path = os.path.abspath(output_path)
        with tempfile.TemporaryDirectory(prefix='excel-to-pdf-') as temp_dir:
            source_path = os.path.abspath(input_path)
            source_extension = os.path.splitext(source_path)[1].lower()

            if source_extension == '.xls':
                source_path = self._run_libreoffice_convert(
                    source_path,
                    'xlsx',
                    temp_dir,
                    os.path.join(temp_dir, 'profile-xls'),
                )

            prepared_extension = '.xlsm' if source_extension == '.xlsm' else '.xlsx'
            prepared_path = os.path.join(temp_dir, f'prepared{prepared_extension}')
            self._prepare_libreoffice_workbook(source_path, prepared_path, options)
            generated_pdf = self._run_libreoffice_convert(
                prepared_path,
                'pdf',
                temp_dir,
                os.path.join(temp_dir, 'profile-pdf'),
            )

            if os.path.exists(output_path):
                os.remove(output_path)
            shutil.move(generated_pdf, output_path)

        return {'method': 'libreoffice'}

    def _get_html_fallback_options(self, input_path: str, options: Dict[str, Any]) -> Dict[str, Any]:
        fallback_options = dict(options)
        orientation = str(options.get('excel_orientation') or 'auto').lower()
        if orientation == 'auto':
            try:
                from openpyxl import load_workbook

                workbook = load_workbook(input_path, read_only=True, data_only=True)
                try:
                    max_columns = max((sheet.max_column or 0) for sheet in workbook.worksheets)
                finally:
                    workbook.close()
                orientation = 'landscape' if max_columns > 6 else 'portrait'
            except Exception:
                orientation = 'landscape'

        fallback_options['orientation'] = orientation
        return fallback_options

    def convert(self, input_path: str, output_path: str, **options) -> Dict[str, Any]:
        """执行转换（混合策略：优先Excel COM，降级LibreOffice）"""
        self.validate_input(input_path)
        self.update_progress(input_path, 5)
        
        error_messages = []
        
        # 策略1: Microsoft Excel (优化版)
        try:
            self.update_progress(input_path, 20)
            result = self._convert_with_excel(input_path, output_path, **options)
            self.update_progress(input_path, 100)
            return {
                'success': True,
                'output_path': output_path,
                'method': result['method'],
                'size': self.get_output_size(output_path)
            }
        except Exception as e:
            error_messages.append(f"Microsoft Excel failed: {str(e)}")
            print(f"[ExcelToPdf] Excel conversion failed: {e}")
        
        # 策略2: LibreOffice
        if self.soffice_path:
            try:
                self.update_progress(input_path, 50)
                result = self._convert_with_libreoffice(input_path, output_path, **options)
                self.update_progress(input_path, 100)
                return {
                    'success': True,
                    'output_path': output_path,
                    'method': result['method'],
                    'size': self.get_output_size(output_path)
                }
            except Exception as e:
                error_messages.append(f"LibreOffice failed: {str(e)}")
                print(f"[ExcelToPdf] LibreOffice conversion failed: {e}")
        else:
            error_messages.append("LibreOffice not available")
            print("[ExcelToPdf] LibreOffice not available")
            
        # 策略3: HTML 中转 (Excel -> HTML -> PDF via browser)
        try:
            self.update_progress(input_path, 60)
            print("[ExcelToPdf] Trying HTML-based fallback...")
            import tempfile
            from .excel_to_html import ExcelToHtmlConverter
            from .html_to_pdf import HtmlToPdfConverter
            
            html_converter = ExcelToHtmlConverter()
            pdf_converter = HtmlToPdfConverter()
            
            with tempfile.NamedTemporaryFile(suffix='.html', delete=False) as tmp:
                html_path = tmp.name
            fallback_options = self._get_html_fallback_options(input_path, options)
            html_result = html_converter.convert(input_path, html_path, **fallback_options)
            if not html_result.get('success'):
                raise Exception(html_result.get('error') or 'Excel to HTML conversion failed')
            pdf_result = pdf_converter.convert(html_path, output_path, **fallback_options)
            os.unlink(html_path)
            
            self.update_progress(input_path, 100)
            return {
                'success': True,
                'output_path': output_path,
                'method': 'html_fallback',
                'size': self.get_output_size(output_path)
            }
        except Exception as e:
            error_messages.append(f"HTML fallback failed: {str(e)}")
            print(f"[ExcelToPdf] HTML fallback failed: {e}")
            
        # 所有策略都失败
        return {
            'success': False,
            'error': "All conversion strategies failed.\n" + "\n".join(error_messages)
        }
