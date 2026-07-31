from io import BytesIO
from typing import Any, Dict, List, Tuple

from docx import Document
from docx.table import Table
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from .base import BaseConverter
from backend.utils.docx_content import docx_paragraph_images, docx_table_rows, iter_docx_blocks
from backend.utils.pptx_tables import add_paginated_table_slides


EMU_PER_INCH = 914400


class DocxToPptConverter(BaseConverter):
    """Word 转 PPT 转换器"""

    def __init__(self):
        super().__init__()

    def _slide_size_inches(self, presentation: Presentation) -> Tuple[float, float]:
        return presentation.slide_width / EMU_PER_INCH, presentation.slide_height / EMU_PER_INCH

    def _parse_color(self, value: str, fallback: str = '#c8c8c8') -> RGBColor:
        color = str(value or fallback).strip().lstrip('#')
        if len(color) == 3:
            color = ''.join(char * 2 for char in color)
        if len(color) != 6:
            color = fallback.lstrip('#')
        try:
            return RGBColor(int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16))
        except Exception:
            color = fallback.lstrip('#')
            return RGBColor(int(color[0:2], 16), int(color[2:4], 16), int(color[4:6], 16))

    def _add_watermark(self, presentation: Presentation, slide, options: Dict[str, Any]) -> None:
        text = str(options.get('watermark_text') or '').strip()
        if not text:
            return

        slide_width, slide_height = self._slide_size_inches(presentation)
        box = slide.shapes.add_textbox(
            Inches(0.9),
            Inches(max(0.4, slide_height / 2 - 0.45)),
            Inches(max(1.0, slide_width - 1.8)),
            Inches(0.9),
        )
        box.rotation = int(options.get('watermark_angle') or 35)
        frame = box.text_frame
        frame.clear()
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = text
        run.font.size = Pt(max(18, int(options.get('watermark_size') or 42)))
        run.font.bold = False
        run.font.color.rgb = self._parse_color(str(options.get('watermark_color') or '#c8c8c8'))

    def _add_title(self, slide, text: str, width: float) -> None:
        title_box = slide.shapes.add_textbox(Inches(0.55), Inches(0.24), Inches(width - 1.1), Inches(0.55))
        frame = title_box.text_frame
        frame.clear()
        paragraph = frame.paragraphs[0]
        paragraph.text = text[:96] if text else 'Document'
        paragraph.font.size = Pt(19)
        paragraph.font.bold = False

    def _add_text_box(self, slide, paragraphs: List[str], left: float, top: float, width: float, height: float) -> None:
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        frame = box.text_frame
        frame.clear()
        frame.word_wrap = True
        total_chars = sum(len(item) for item in paragraphs)
        font_size = 17 if total_chars < 520 else (14 if total_chars < 950 else 12)
        for index, item in enumerate(paragraphs or ['']):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = item
            paragraph.space_after = Pt(6)
            paragraph.font.size = Pt(font_size)

    def _add_images(self, slide, images: List[dict], left: float, top: float, width: float, height: float) -> None:
        if not images:
            return

        max_images = min(len(images), 3)
        slot_height = height / max_images
        for index, image_info in enumerate(images[:max_images]):
            blob = image_info.get('blob')
            if not blob:
                continue
            try:
                picture = slide.shapes.add_picture(
                    BytesIO(blob),
                    Inches(left),
                    Inches(top + index * slot_height),
                    width=Inches(width),
                )
                max_height = Inches(max(0.5, slot_height - 0.15))
                if picture.height > max_height:
                    scale = max_height / picture.height
                    picture.width = int(picture.width * scale)
                    picture.height = int(picture.height * scale)
                if picture.width > Inches(width):
                    scale = Inches(width) / picture.width
                    picture.width = int(picture.width * scale)
                    picture.height = int(picture.height * scale)
            except Exception:
                continue

    def _add_content_slide(
        self,
        presentation: Presentation,
        paragraphs: List[str],
        images: List[dict],
        options: Dict[str, Any],
    ) -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide_width, slide_height = self._slide_size_inches(presentation)
        title = paragraphs[0] if paragraphs else 'Document image'
        body = paragraphs[1:] if len(paragraphs) > 1 else paragraphs
        self._add_title(slide, title, slide_width)

        content_top = 0.95
        content_height = slide_height - 1.35
        if images:
            text_width = max(3.6, slide_width * 0.48)
            image_left = text_width + 0.85
            image_width = max(3.0, slide_width - image_left - 0.55)
            self._add_text_box(slide, body or paragraphs, 0.55, content_top, text_width, content_height)
            self._add_images(slide, images, image_left, content_top, image_width, content_height)
        else:
            self._add_text_box(slide, body or paragraphs, 0.65, content_top, slide_width - 1.3, content_height)
        self._add_watermark(presentation, slide, options)

    def _flush_text_slides(
        self,
        presentation: Presentation,
        paragraphs: List[str],
        options: Dict[str, Any],
    ) -> None:
        if not paragraphs:
            return
        chunk: List[str] = []
        char_count = 0
        for paragraph in paragraphs:
            if chunk and char_count + len(paragraph) > 1050:
                self._add_content_slide(presentation, chunk, [], options)
                chunk = []
                char_count = 0
            chunk.append(paragraph)
            char_count += len(paragraph)
        if chunk:
            self._add_content_slide(presentation, chunk, [], options)

    def convert(self, input_path: str, output_path: str, **options) -> Dict[str, Any]:
        self.validate_input(input_path)
        self.update_progress(input_path, 5)

        document = Document(input_path)
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)

        pending_paragraphs: List[str] = []
        for block in iter_docx_blocks(document):
            if isinstance(block, Table):
                self._flush_text_slides(presentation, pending_paragraphs, options)
                pending_paragraphs = []
                before_count = len(presentation.slides)
                add_paginated_table_slides(presentation, 'Document table', docx_table_rows(block))
                for index in range(before_count, len(presentation.slides)):
                    self._add_watermark(presentation, presentation.slides[index], options)
                continue

            text = block.text.strip()
            images = docx_paragraph_images(block)
            if images:
                self._flush_text_slides(presentation, pending_paragraphs, options)
                pending_paragraphs = []
                self._add_content_slide(presentation, [text] if text else [], images, options)
                continue
            if text:
                pending_paragraphs.append(text)

        self._flush_text_slides(presentation, pending_paragraphs, options)

        if not presentation.slides:
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            self._add_watermark(presentation, slide, options)
        presentation.save(output_path)
        self.update_progress(input_path, 100)
        return {
            'success': True,
            'output_path': output_path,
            'size': self.get_output_size(output_path),
            'slides': len(presentation.slides),
        }
