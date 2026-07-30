from typing import Dict, Any
from pptx import Presentation
from pptx.util import Inches
from .base import BaseConverter
from backend.utils.pptx_tables import add_paginated_table_slides
from backend.utils.spreadsheet_data import read_spreadsheet_rows

class ExcelToPptConverter(BaseConverter):
    """Excel 转 PPT 转换器"""
    
    def __init__(self):
        super().__init__()

    def convert(self, input_path: str, output_path: str, **options) -> Dict[str, Any]:
        self.validate_input(input_path)
        self.update_progress(input_path, 5)
        
        presentation = Presentation()
        presentation.slide_width = Inches(10)
        presentation.slide_height = Inches(7.5)
        with read_spreadsheet_rows(input_path) as sheets:
            for sheet_name, sheet_rows in sheets:
                rows = [
                    ['' if value is None else str(value) for value in row]
                    for row in sheet_rows
                ]
                while rows and not any(value for value in rows[-1]):
                    rows.pop()
                add_paginated_table_slides(presentation, sheet_name, rows)

        if not presentation.slides:
            presentation.slides.add_slide(presentation.slide_layouts[6])
        presentation.save(output_path)
        self.update_progress(input_path, 100)
        return {
            'success': True,
            'output_path': output_path,
            'size': self.get_output_size(output_path),
            'slides': len(presentation.slides),
        }
