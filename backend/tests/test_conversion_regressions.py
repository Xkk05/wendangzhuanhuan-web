import base64
import zipfile
from io import BytesIO
from pathlib import Path

import fitz
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.shared import Inches as DocxInches
from openpyxl import Workbook
from PIL import Image, ImageChops, ImageStat
from pptx import Presentation

from backend.converters.docx_to_epub import DocxToEpubConverter
from backend.converters.docx_to_image import DocxToImageConverter
from backend.converters.docx_to_pdf import DocxToPdfConverter
from backend.converters.docx_to_ppt import DocxToPptConverter
from backend.converters.excel_to_pdf import ExcelToPdfConverter
from backend.converters.excel_to_html import ExcelToHtmlConverter
from backend.converters.excel_to_ppt import ExcelToPptConverter
from backend.converters.html_to_gif import HtmlToGifConverter
from backend.converters.html_to_docx import HtmlToDocxConverter
from backend.converters.html_to_image import HtmlToImageConverter
from backend.converters.html_to_json import HtmlToJsonConverter
from backend.converters.html_to_markdown import HtmlToMarkdownConverter
from backend.converters.html_to_pdf import HtmlToPdfConverter
from backend.converters.html_to_svg import HtmlToSvgConverter
from backend.converters.html_to_txt import HtmlToTxtConverter
from backend.converters.json_to_html import JsonToHtmlConverter
from backend.converters.json_to_image import JsonToImageConverter
from backend.converters.json_to_pdf import JsonToPdfConverter
from backend.converters.json_to_svg import JsonToSvgConverter
from backend.converters.pdf_to_md import PdfToMdConverter
from backend.converters.pdf_to_docx import PdfToDocxConverter
from backend.converters.pdf_to_ppt import PdfToPptConverter
from backend.converters.txt_to_image import TxtToImageConverter
from backend.converters.txt_to_pdf import TxtToPdfConverter
from backend.converters.txt_to_speech import TxtToSpeechConverter
from backend.converters.xml_to_html import XmlToHtmlConverter
from backend.converters.xml_to_image import XmlToImageConverter
from backend.converters.xml_to_pdf import XmlToPdfConverter
from backend.converters.xml_to_svg import XmlToSvgConverter
from backend.services.converter_service import ConverterService
from backend.utils.text_utils import detect_tts_language, normalize_tts_language
from backend.utils.spreadsheet_data import read_spreadsheet_rows


def _assert_image_has_content(path: Path) -> Image.Image:
    image = Image.open(path).convert('RGB')
    background = Image.new('RGB', image.size, image.getpixel((0, 0)))
    difference = ImageChops.difference(image, background).convert('L')
    assert ImageStat.Stat(difference).extrema[0][1] > 20
    return image


def _assert_valid_pptx(path: Path, minimum_slides: int = 1) -> None:
    assert path.exists() and path.stat().st_size > 0
    with zipfile.ZipFile(path) as archive:
        slide_names = [
            name for name in archive.namelist()
            if name.startswith('ppt/slides/slide') and name.endswith('.xml')
        ]
        assert len(slide_names) >= minimum_slides


def _make_text_pdf(path: Path, text: str = 'Regression PDF content') -> None:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    document.save(path)
    document.close()


def _make_docx(path: Path) -> None:
    document = Document()
    document.add_heading('Watermark regression', level=1)
    document.add_paragraph('DOCX image conversion content')
    document.save(path)


def _make_xlsx(path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Regression'
    worksheet.append(['Name', 'January', 'February', 'March'])
    for index in range(1, 8):
        worksheet.append([f'Item {index}', index * 10, index * 20, index * 30])
    workbook.save(path)


def _presentation_text(path: Path) -> str:
    presentation = Presentation(path)
    values = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, 'has_text_frame', False):
                values.append(shape.text)
            if getattr(shape, 'has_table', False):
                for row in shape.table.rows:
                    values.extend(cell.text for cell in row.cells)
    return '\n'.join(values)


def test_docx_png_and_jpg_apply_watermark(tmp_path):
    source = tmp_path / 'source.docx'
    _make_docx(source)
    converter = DocxToImageConverter()

    for extension in ('png', 'jpg'):
        plain = tmp_path / f'plain.{extension}'
        marked = tmp_path / f'marked.{extension}'
        converter.convert(str(source), str(plain), quality=90)
        converter.convert(
            str(source),
            str(marked),
            quality=90,
            watermark_text='Watermark 2026',
            watermark_opacity=90,
            watermark_size=44,
            watermark_color='#335577',
            watermark_angle=0,
            watermark_position='center',
        )

        plain_image = Image.open(plain).convert('RGB')
        marked_image = Image.open(marked).convert('RGB')
        assert plain_image.size == marked_image.size
        difference = ImageChops.difference(plain_image, marked_image).convert('L')
        assert ImageStat.Stat(difference).mean[0] > 0.05


def test_docx_portrait_pdf_preprocess_fits_wide_tables(tmp_path):
    source = tmp_path / 'wide-table.docx'
    output = tmp_path / 'wide-table.pdf'
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    table = document.add_table(rows=2, cols=4)
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.width = DocxInches(3)
            cell.text = f'R{row_index + 1}C{col_index + 1} 宽表格内容'
    document.save(source)

    prepared_path, cleanup_path = DocxToPdfConverter()._prepare_docx_orientation_for_libreoffice(
        str(source),
        str(output),
        'portrait',
    )

    try:
        assert cleanup_path == prepared_path
        prepared = Document(prepared_path)
        prepared_section = prepared.sections[0]
        assert prepared_section.orientation == WD_ORIENT.PORTRAIT
        assert prepared_section.page_width < prepared_section.page_height

        prepared_table = prepared.tables[0]
        tbl_pr = prepared_table._tbl.tblPr
        tbl_w = tbl_pr.find(qn('w:tblW'))
        tbl_layout = tbl_pr.find(qn('w:tblLayout'))
        assert tbl_w is not None
        assert tbl_w.get(qn('w:type')) == 'pct'
        assert tbl_w.get(qn('w:w')) == '5000'
        assert tbl_layout is not None
        assert tbl_layout.get(qn('w:type')) == 'autofit'

        available_width = (
            prepared_section.page_width
            - prepared_section.left_margin
            - prepared_section.right_margin
        )
        first_row_width = sum(
            int(cell.width or 0)
            for cell in prepared_table.rows[0].cells
        )
        assert first_row_width <= available_width
    finally:
        if cleanup_path and Path(cleanup_path).exists():
            Path(cleanup_path).unlink()


def test_html_pdf_print_css_fits_wide_content():
    html = HtmlToPdfConverter()._prepare_html_content(
        '<html><head></head><body>'
        '<table style="width: 1800px"><tr><td>很长的表格内容</td></tr></table>'
        '<img style="width: 1600px" src="sample.png">'
        '<pre>https://example.com/very/long/path/that/should/wrap</pre>'
        '</body></html>',
        {},
    )

    assert 'table-layout: fixed !important' in html
    assert 'body *' in html
    assert 'max-width: 100% !important' in html
    assert 'overflow-wrap: anywhere' in html
    assert 'white-space: pre-wrap' in html


def test_txt_pdf_html_keeps_encoding_spaces_and_wrapping(tmp_path, monkeypatch):
    source = tmp_path / 'source.txt'
    output = tmp_path / 'source.pdf'
    source.write_bytes('第一列    第二列\n超长链接 https://example.com/'.encode('gb18030'))
    captured = {}

    def fake_convert(self, input_path, output_path, **options):
        captured['html'] = Path(input_path).read_text(encoding='utf-8')
        Path(output_path).write_bytes(b'%PDF-test')
        return {'success': True, 'output_path': output_path, 'size': Path(output_path).stat().st_size}

    monkeypatch.setattr(HtmlToPdfConverter, 'convert', fake_convert)

    result = TxtToPdfConverter().convert(str(source), str(output))

    assert result['success']
    assert '第一列    第二列' in captured['html']
    assert '&nbsp;' not in captured['html']
    assert 'overflow-wrap: anywhere' in captured['html']


def test_html_code_pdf_keeps_chinese_after_cleanup(tmp_path):
    source = tmp_path / 'source.html'
    output = tmp_path / 'source.pdf'
    source.write_bytes(
        '<html><head><meta charset="gb18030"><style>p{color:red}</style></head>'
        '<body><p>中文内容正常显示</p><script>window.bad=true</script></body></html>'.encode('gb18030')
    )

    result = HtmlToPdfConverter().convert(
        str(source),
        str(output),
        code_mode=True,
        encoding='gb18030',
        css_handling='remove',
        remove_scripts=True,
    )

    assert result['success']
    with fitz.open(output) as document:
        text = ''.join(page.get_text() for page in document)
    assert '中文内容正常显示' in text
    assert 'window.bad' not in text


def test_html_svg_and_image_formats_generate_valid_outputs(tmp_path):
    source = tmp_path / 'source.html'
    source.write_text(
        '<!doctype html><meta charset="utf-8"><h1>转换标题</h1><p>中文正文内容</p>',
        encoding='utf-8',
    )

    svg = tmp_path / 'output.svg'
    assert HtmlToSvgConverter().convert(str(source), str(svg))['success']
    svg_text = svg.read_text(encoding='utf-8')
    assert '<svg' in svg_text and ('中文正文内容' in svg_text or 'data:image/png;base64,' in svg_text)

    dimensions = []
    for extension, converter in (
        ('png', HtmlToImageConverter()),
        ('jpg', HtmlToImageConverter()),
        ('gif', HtmlToGifConverter()),
    ):
        output = tmp_path / f'output.{extension}'
        assert converter.convert(str(source), str(output))['success']
        image = _assert_image_has_content(output)
        dimensions.append(image.size)
    assert dimensions[0] == dimensions[1] == dimensions[2]


def test_html_image_trims_only_trailing_blank_pdf_pages(tmp_path):
    pdf = tmp_path / 'with-blank-tail.pdf'
    document = fitz.open()
    content_page = document.new_page()
    content_page.insert_text((72, 72), 'Keep this content page')
    document.new_page()
    document.new_page()
    document.save(pdf)
    document.close()

    removed = HtmlToImageConverter()._trim_trailing_blank_pages(str(pdf))

    assert removed == 2
    with fitz.open(pdf) as trimmed:
        assert len(trimmed) == 1
        assert 'Keep this content page' in trimmed[0].get_text()


def test_pdf_markdown_and_powerpoint_outputs_are_usable(tmp_path):
    source = tmp_path / 'source.pdf'
    markdown = tmp_path / 'source.md'
    powerpoint = tmp_path / 'source.pptx'
    _make_text_pdf(source)

    md_result = PdfToMdConverter().convert(str(source), str(markdown))
    assert md_result['success']
    assert 'Regression PDF content' in markdown.read_text(encoding='utf-8')

    ppt_result = PdfToPptConverter().convert(str(source), str(powerpoint))
    assert ppt_result['success']
    _assert_valid_pptx(powerpoint)


def test_excel_powerpoint_and_landscape_pdf_outputs_are_usable(tmp_path):
    source = tmp_path / 'source.xlsx'
    _make_xlsx(source)
    powerpoint = tmp_path / 'source.pptx'
    pdf = tmp_path / 'source.pdf'

    ppt_result = ExcelToPptConverter().convert(str(source), str(powerpoint))
    assert ppt_result['success']
    _assert_valid_pptx(powerpoint)

    pdf_result = ExcelToPdfConverter().convert(
        str(source),
        str(pdf),
        page_size='A4',
        excel_orientation='landscape',
        excel_scale_mode='fit_width',
    )
    assert pdf_result['success']
    with fitz.open(pdf) as document:
        assert len(document) >= 1
        assert document[0].rect.width > document[0].rect.height


def test_txt_chinese_images_and_speech_language(tmp_path, monkeypatch):
    source = tmp_path / 'source.txt'
    source.write_text('这是中文语音和图片回归测试。', encoding='utf-8')

    for extension in ('png', 'jpg'):
        output = tmp_path / f'output.{extension}'
        assert TxtToImageConverter().convert(str(source), str(output))['success']
        _assert_image_has_content(output)

    assert detect_tts_language(source.read_text(encoding='utf-8')) == 'zh-CN'
    captured = {}

    def fake_gtts(self, text, output_path, language, rate, pitch):
        captured.update(text=text, language=language, rate=rate, pitch=pitch)
        Path(output_path).write_bytes(b'ID3-regression-audio')
        return {'success': True, 'output_path': output_path, 'method': 'test-gtts', 'language': language}

    monkeypatch.setattr(TxtToSpeechConverter, '_convert_with_gtts', fake_gtts)
    audio = tmp_path / 'output.mp3'
    result = TxtToSpeechConverter().convert(str(source), str(audio))
    assert result['success']
    assert captured['language'] == 'zh-CN'
    assert captured['text'].startswith('这是中文')


def test_txt_crlf_and_lf_render_identically(tmp_path):
    lf_source = tmp_path / 'lf.txt'
    crlf_source = tmp_path / 'crlf.txt'
    lf_source.write_bytes('第一行\n第二行\n\n第四行'.encode('utf-8'))
    crlf_source.write_bytes('第一行\r\n第二行\r\n\r\n第四行'.encode('utf-8'))

    lf_output = tmp_path / 'lf.png'
    crlf_output = tmp_path / 'crlf.png'
    converter = TxtToImageConverter()
    converter.convert(str(lf_source), str(lf_output))
    converter.convert(str(crlf_source), str(crlf_output))

    assert ImageChops.difference(Image.open(lf_output), Image.open(crlf_output)).getbbox() is None


def test_tts_language_detection_and_normalization_cover_advertised_scripts():
    assert detect_tts_language('これは日本語です。') == 'ja'
    assert detect_tts_language('한국어 음성 테스트입니다.') == 'ko'
    assert detect_tts_language('هذا اختبار صوتي باللغة العربية.') == 'ar'
    assert detect_tts_language('Это проверка русского голоса.') == 'ru'
    assert normalize_tts_language('pt-BR') == 'pt'
    assert normalize_tts_language('zh_TW') == 'zh-TW'
    assert normalize_tts_language('unsupported', text='这是中文。') == 'zh-CN'


def test_xml_display_outputs_remove_tags_and_keep_all_text(tmp_path):
    source = tmp_path / 'source.xml'
    source.write_text(
        '<document><title>XML标题</title><para>第一段正文</para><para>第二段正文</para></document>',
        encoding='utf-8',
    )

    html_output = tmp_path / 'output.html'
    assert XmlToHtmlConverter().convert(str(source), str(html_output))['success']
    html_text = html_output.read_text(encoding='utf-8')
    assert '<para>' not in html_text and '&lt;para&gt;' not in html_text

    svg_output = tmp_path / 'output.svg'
    assert XmlToSvgConverter().convert(str(source), str(svg_output))['success']
    svg_text = svg_output.read_text(encoding='utf-8')
    assert '第一段正文' in svg_text and '第二段正文' in svg_text

    pdf_output = tmp_path / 'output.pdf'
    assert XmlToPdfConverter().convert(str(source), str(pdf_output))['success']
    with fitz.open(pdf_output) as document:
        pdf_text = ''.join(page.get_text() for page in document)
    assert '第一段正文' in pdf_text and '第二段正文' in pdf_text

    for extension in ('png', 'jpg'):
        image_output = tmp_path / f'output.{extension}'
        assert XmlToImageConverter().convert(str(source), str(image_output))['success']
        _assert_image_has_content(image_output)


def test_json_display_outputs_remove_structure_and_keep_text(tmp_path):
    source = tmp_path / 'source.json'
    source.write_text(
        '{"title":"JSON标题","node_type":"paragraph","content":"正文内容","items":["第一项","第二项"]}',
        encoding='utf-8',
    )

    html_output = tmp_path / 'output.html'
    assert JsonToHtmlConverter().convert(str(source), str(html_output))['success']
    html_text = html_output.read_text(encoding='utf-8')
    assert 'node_type' not in html_text and '&quot;content&quot;' not in html_text

    svg_output = tmp_path / 'output.svg'
    assert JsonToSvgConverter().convert(str(source), str(svg_output))['success']
    svg_text = svg_output.read_text(encoding='utf-8')
    assert '正文内容' in svg_text and 'node_type' not in svg_text

    pdf_output = tmp_path / 'output.pdf'
    assert JsonToPdfConverter().convert(str(source), str(pdf_output))['success']
    with fitz.open(pdf_output) as document:
        pdf_text = ''.join(page.get_text() for page in document)
    assert '正文内容' in pdf_text and 'node_type' not in pdf_text

    for extension in ('png', 'jpg'):
        image_output = tmp_path / f'output.{extension}'
        assert JsonToImageConverter().convert(str(source), str(image_output))['success']
        _assert_image_has_content(image_output)


def test_office_exports_preserve_tables_and_wide_sheet_values(tmp_path):
    docx_source = tmp_path / 'table.docx'
    document = Document()
    document.add_heading('交付文档', level=1)
    table = document.add_table(rows=2, cols=3)
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            cell.text = f'文档表格-{row_index}-{col_index}'
    document.save(docx_source)

    ppt_output = tmp_path / 'table.pptx'
    epub_output = tmp_path / 'table.epub'
    assert DocxToPptConverter().convert(str(docx_source), str(ppt_output))['success']
    assert DocxToEpubConverter().convert(str(docx_source), str(epub_output))['success']
    ppt_text = _presentation_text(ppt_output)
    assert '文档表格-0-0' in ppt_text and '文档表格-1-2' in ppt_text
    with zipfile.ZipFile(epub_output) as archive:
        epub_text = '\n'.join(
            archive.read(name).decode('utf-8')
            for name in archive.namelist()
            if name.endswith('.xhtml')
        )
    assert '<table' in epub_text and '文档表格-1-2' in epub_text

    excel_source = tmp_path / 'wide.xlsx'
    workbook = Workbook()
    worksheet = workbook.active
    for row in range(1, 22):
        worksheet.append([f'R{row}C{column}' for column in range(1, 13)])
    workbook.save(excel_source)
    excel_ppt = tmp_path / 'wide.pptx'
    assert ExcelToPptConverter().convert(str(excel_source), str(excel_ppt))['success']
    excel_ppt_text = _presentation_text(excel_ppt)
    assert 'R1C12' in excel_ppt_text
    assert 'R21C1' in excel_ppt_text
    assert 'R21C12' in excel_ppt_text


def test_docx_powerpoint_keeps_images_and_watermark(tmp_path):
    docx_source = tmp_path / 'image.docx'
    image_source = tmp_path / 'inline.png'
    Image.new('RGB', (220, 120), (60, 120, 180)).save(image_source)

    document = Document()
    document.add_heading('图文演示', level=1)
    document.add_paragraph('这是一段需要保留的正文。')
    document.add_picture(str(image_source), width=DocxInches(1.6))
    document.save(docx_source)

    ppt_output = tmp_path / 'image.pptx'
    assert DocxToPptConverter().convert(
        str(docx_source),
        str(ppt_output),
        watermark_text='内部水印',
        watermark_angle=0,
    )['success']

    ppt_text = _presentation_text(ppt_output)
    assert '图文演示' in ppt_text
    assert '这是一段需要保留的正文' in ppt_text
    assert '内部水印' in ppt_text
    with zipfile.ZipFile(ppt_output) as archive:
        assert any(name.startswith('ppt/media/image') for name in archive.namelist())


def test_html_text_exports_preserve_tables_and_exclude_hidden_content(tmp_path):
    source = tmp_path / 'structured.html'
    source.write_text(
        '<!doctype html><html><head><meta charset="utf-8">'
        '<style>.secret { display: none; }</style></head><body>'
        '<!--不应进入结果的注释--><h1>可见标题</h1><p>可见正文</p>'
        '<p class="secret">类隐藏内容</p><p style="visibility:hidden">内联隐藏内容</p>'
        '<table><thead><tr><th>姓名</th><th>分数</th></tr></thead>'
        '<tbody><tr><td>小明</td><td>42</td></tr></tbody></table></body></html>',
        encoding='utf-8',
    )

    markdown = tmp_path / 'output.md'
    txt = tmp_path / 'output.txt'
    docx = tmp_path / 'output.docx'
    json_output = tmp_path / 'output.json'
    assert HtmlToMarkdownConverter().convert(str(source), str(markdown))['success']
    assert HtmlToTxtConverter().convert(str(source), str(txt))['success']
    assert HtmlToDocxConverter().convert(str(source), str(docx))['success']
    assert HtmlToJsonConverter().convert(str(source), str(json_output))['success']

    markdown_text = markdown.read_text(encoding='utf-8')
    assert '| 姓名 | 分数 |' in markdown_text
    for output_text in (markdown_text, txt.read_text(encoding='utf-8'), json_output.read_text(encoding='utf-8')):
        assert '可见正文' in output_text
        assert '类隐藏内容' not in output_text
        assert '内联隐藏内容' not in output_text
        assert '不应进入结果的注释' not in output_text

    converted_docx = Document(docx)
    assert converted_docx.tables
    assert converted_docx.tables[0].cell(1, 1).text == '42'
    docx_text = '\n'.join(paragraph.text for paragraph in converted_docx.paragraphs)
    assert '类隐藏内容' not in docx_text and '内联隐藏内容' not in docx_text


def test_html_docx_preserves_nested_blocks_without_duplicate_list_text(tmp_path):
    source = tmp_path / 'nested.html'
    source.write_text(
        '<html><body><main><section><h2>嵌套标题</h2><p>嵌套正文</p>'
        '<table><tr><th>项目</th><th>值</th></tr><tr><td>甲</td><td>42</td></tr></table>'
        '<ul><li>一级<ul><li>二级</li></ul></li></ul></section></main></body></html>',
        encoding='utf-8',
    )
    output = tmp_path / 'nested.docx'

    assert HtmlToDocxConverter().convert(str(source), str(output))['success']

    document = Document(output)
    paragraph_text = [paragraph.text for paragraph in document.paragraphs]
    assert paragraph_text.count('嵌套标题') == 1
    assert paragraph_text.count('嵌套正文') == 1
    assert paragraph_text.count('一级') == 1
    assert paragraph_text.count('二级') == 1
    assert len(document.tables) == 1
    assert document.tables[0].cell(1, 1).text == '42'


def test_html_docx_embeds_local_and_data_uri_images(tmp_path):
    local_image = tmp_path / 'local.png'
    Image.new('RGB', (160, 90), (80, 140, 200)).save(local_image)

    inline_buffer = BytesIO()
    Image.new('RGB', (120, 80), (180, 90, 120)).save(inline_buffer, format='PNG')
    inline_data = base64.b64encode(inline_buffer.getvalue()).decode('ascii')

    source = tmp_path / 'images.html'
    source.write_text(
        '<!doctype html><html><body><h1>图文页面</h1><p>图片前正文</p>'
        '<img src="local.png" alt="本地图片">'
        f'<p>内联图片<img src="data:image/png;base64,{inline_data}" alt="内联图片"></p>'
        '</body></html>',
        encoding='utf-8',
    )
    output = tmp_path / 'images.docx'

    assert HtmlToDocxConverter().convert(str(source), str(output))['success']

    document = Document(output)
    assert len(document.inline_shapes) == 2
    text = '\n'.join(paragraph.text for paragraph in document.paragraphs)
    assert '图文页面' in text and '图片前正文' in text
    with zipfile.ZipFile(output) as archive:
        media_files = [name for name in archive.namelist() if name.startswith('word/media/image')]
    assert len(media_files) == 2


def test_excel_html_escapes_sheet_names_and_cell_content(tmp_path):
    source = tmp_path / 'special.xlsx'
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = '研发 & 测试'
    worksheet.append(['名称', '<script>alert("x")</script>'])
    workbook.save(source)
    output = tmp_path / 'special.html'

    result = ExcelToHtmlConverter().convert(str(source), str(output))

    assert result['success']
    html = output.read_text(encoding='utf-8')
    assert '研发 &amp; 测试' in html
    assert '&lt;script&gt;alert(&quot;x&quot;)&lt;/script&gt;' in html
    assert '<script>alert("x")</script>' not in html


def test_legacy_xls_reader_uses_xlrd_values(tmp_path, monkeypatch):
    import xlrd

    source = tmp_path / 'legacy.xls'
    source.write_bytes(b'legacy-fixture')

    class FakeCell:
        def __init__(self, value, cell_type=xlrd.XL_CELL_TEXT):
            self.value = value
            self.ctype = cell_type

    class FakeSheet:
        name = '旧版表格'
        nrows = 2
        ncols = 2
        values = [
            [FakeCell('名称'), FakeCell('启用')],
            [FakeCell('项目甲'), FakeCell(1, xlrd.XL_CELL_BOOLEAN)],
        ]

        def cell(self, row_index, column_index):
            return self.values[row_index][column_index]

    class FakeBook:
        datemode = 0

        def __init__(self):
            self.released = False

        def sheets(self):
            return [FakeSheet()]

        def release_resources(self):
            self.released = True

    fake_book = FakeBook()
    monkeypatch.setattr(xlrd, 'open_workbook', lambda *_args, **_kwargs: fake_book)

    with read_spreadsheet_rows(str(source)) as sheets:
        assert sheets == [('旧版表格', [['名称', '启用'], ['项目甲', True]])]
    assert fake_book.released


def test_pdf_doc_alias_uses_docx_extension(tmp_path, monkeypatch):
    source = tmp_path / 'source.pdf'
    source.write_bytes(b'%PDF-test')
    service = object.__new__(ConverterService)
    service.converters = {('pdf', 'doc'): PdfToDocxConverter()}

    def fake_convert(_input_path, output_path, **_options):
        Path(output_path).write_bytes(b'PK\x03\x04-docx')
        return {'success': True, 'output_path': output_path}

    monkeypatch.setattr(service.converters[('pdf', 'doc')], 'convert', fake_convert)
    result = service.convert_file(str(source), 'doc', original_filename='报告.pdf')

    assert result['display_name'].endswith('.docx')
    assert result['filename'].endswith('.docx')
