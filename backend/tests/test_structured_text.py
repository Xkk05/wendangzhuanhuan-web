import base64
import csv
import json

import pytest
from fastapi import HTTPException
from PIL import Image
from openpyxl import Workbook, load_workbook

from backend.api.routes import GUEST_TOKEN, _assert_processing_access
from backend.converters.excel_to_pdf import ExcelToPdfConverter
from backend.converters.html_to_pdf import HtmlToPdfConverter
from backend.converters.excel_to_html import ExcelToHtmlConverter
from backend.converters.json_to_base64 import JsonToBase64Converter
from backend.converters.json_to_html import JsonToHtmlConverter
from backend.converters.json_to_csv import JsonToCsvConverter
from backend.converters.json_to_xml import JsonToXmlConverter
from backend.converters.json_to_xlsx import JsonToXlsxConverter
from backend.converters.json_to_svg import JsonToSvgConverter
from backend.converters.json_to_yaml import JsonToYamlConverter
from backend.converters.pdf_to_image import PdfToImageConverter
from backend.converters.xml_to_html import XmlToHtmlConverter
from backend.converters.xml_to_csv import XmlToCsvConverter
from backend.converters.xml_to_svg import XmlToSvgConverter
from backend.converters.xml_to_txt import XmlToTxtConverter
from backend.converters.xml_to_xlsx import XmlToXlsxConverter
from backend.utils.structured_text import extract_json_texts, extract_xml_texts


def test_extract_xml_texts_returns_content_without_element_names():
    from xml.etree import ElementTree as ET

    root = ET.fromstring('<document><title>标题</title><para>第一段</para><para>第二段</para></document>')

    assert extract_xml_texts(root) == ['标题', '第一段', '第二段']


def test_extract_json_texts_returns_all_content_scalars():
    data = {
        'title': '标题',
        'meta': {'node_type': 'paragraph', 'content': '正文'},
        'items': ['第一项', 42, True, False, None, 0],
    }

    assert extract_json_texts(data) == ['标题', '正文', '第一项', '42', 'true', 'false', 'null', '0']


def test_guest_processing_access_requires_login():
    with pytest.raises(HTTPException) as exc_info:
        _assert_processing_access(GUEST_TOKEN)

    assert exc_info.value.status_code == 401
    assert exc_info.value.detail['code'] == 'login_required'


def test_xml_tabular_exports_preserve_paths_attributes_and_repeated_values(tmp_path):
    source = tmp_path / 'sample.xml'
    source.write_text(
        '<document id="doc-7"><title>完整标题</title><para>第一段</para><para>第二段</para>'
        '<items><item code="A"><name>项目甲</name><score>42</score></item>'
        '<item code="B"><name>项目乙</name><score>0</score></item></items></document>',
        encoding='utf-8',
    )

    csv_output = tmp_path / 'sample.csv'
    xlsx_output = tmp_path / 'sample.xlsx'
    assert XmlToCsvConverter().convert(str(source), str(csv_output))['success']
    assert XmlToXlsxConverter().convert(str(source), str(xlsx_output))['success']

    with csv_output.open(encoding='utf-8-sig', newline='') as handle:
        csv_rows = list(csv.DictReader(handle))
    csv_pairs = {(row['path'], row['value']) for row in csv_rows}

    workbook = load_workbook(xlsx_output, data_only=True)
    try:
        worksheet = workbook.active
        xlsx_pairs = {
            (str(row[0]), str(row[1]))
            for row in worksheet.iter_rows(min_row=2, values_only=True)
        }
    finally:
        workbook.close()

    expected_values = {'doc-7', '完整标题', '第一段', '第二段', 'A', '项目甲', '42', 'B', '项目乙', '0'}
    assert expected_values <= {value for _, value in csv_pairs}
    assert expected_values <= {value for _, value in xlsx_pairs}
    assert any('@id' in path for path, _ in csv_pairs)
    assert any('para[2]' in path for path, _ in csv_pairs)


def test_json_tabular_exports_preserve_metadata_arrays_and_numeric_values(tmp_path):
    source = tmp_path / 'sample.json'
    source.write_text(
        json.dumps({
            'title': '完整标题',
            'description': '元数据正文',
            'count': 2,
            'items': [
                {'name': '项目甲', 'score': 42},
                {'name': '项目乙', 'score': 0},
            ],
        }, ensure_ascii=False),
        encoding='utf-8',
    )

    csv_output = tmp_path / 'sample.csv'
    xlsx_output = tmp_path / 'sample.xlsx'
    assert JsonToCsvConverter().convert(str(source), str(csv_output))['success']
    assert JsonToXlsxConverter().convert(str(source), str(xlsx_output))['success']

    with csv_output.open(encoding='utf-8-sig', newline='') as handle:
        csv_rows = list(csv.DictReader(handle))
    csv_pairs = {(row['path'], row['value']) for row in csv_rows}

    workbook = load_workbook(xlsx_output, data_only=True)
    try:
        worksheet = workbook.active
        xlsx_values = {
            str(row[0])
            for row in worksheet.iter_rows(min_row=2, values_only=True)
            if row and row[0] is not None
        }
    finally:
        workbook.close()

    expected_values = {'完整标题', '元数据正文', '2', '项目甲', '42', '项目乙', '0'}
    assert expected_values <= {value for _, value in csv_pairs}
    assert expected_values <= xlsx_values
    assert ('items[2].score', '0') in csv_pairs


def test_nearly_blank_page_is_ignored():
    image = Image.new('RGB', (1000, 1000), 'white')
    image.putpixel((500, 500), (245, 245, 245))

    assert PdfToImageConverter()._is_image_empty(image)


def test_excel_page_settings_honor_user_selection():
    settings = ExcelToPdfConverter()._resolve_page_settings(
        {'page_size': 'A3', 'excel_orientation': 'landscape', 'excel_scale_mode': 'fit_width'},
        col_count=3,
    )

    assert settings == {'paper_size': 8, 'orientation': 2, 'scale_mode': 'fit_width'}


def test_excel_page_settings_preserve_wide_sheet_auto_scale():
    settings = ExcelToPdfConverter()._resolve_page_settings({}, col_count=20)

    assert settings == {'paper_size': 8, 'orientation': 2, 'scale_mode': 'scaled_85'}


def test_libreoffice_workbook_contains_requested_print_settings(tmp_path):
    source = tmp_path / 'source.xlsx'
    prepared = tmp_path / 'prepared.xlsx'
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(['Name', 'January', 'February'])
    worksheet.append(['Item', 10, 20])
    workbook.save(source)
    workbook.close()

    ExcelToPdfConverter()._prepare_libreoffice_workbook(
        str(source),
        str(prepared),
        {
            'page_size': 'A4',
            'excel_orientation': 'landscape',
            'excel_scale_mode': 'fit_width',
        },
    )

    prepared_workbook = load_workbook(prepared)
    try:
        prepared_sheet = prepared_workbook.active
        assert prepared_sheet.page_setup.paperSize == 9
        assert prepared_sheet.page_setup.orientation == 'landscape'
        assert prepared_sheet.page_setup.fitToWidth == 1
        assert prepared_sheet.page_setup.fitToHeight == 0
        assert prepared_sheet.sheet_properties.pageSetUpPr.fitToPage is True
    finally:
        prepared_workbook.close()


def test_xml_display_exports_keep_text_and_remove_source_tags(tmp_path):
    source = tmp_path / 'sample.xml'
    source.write_text(
        '<document><title>转换标题</title><para>第一段正文</para><para>第二段正文</para></document>',
        encoding='utf-8',
    )

    outputs = {
        'html': (XmlToHtmlConverter(), tmp_path / 'sample.html'),
        'txt': (XmlToTxtConverter(), tmp_path / 'sample.txt'),
        'svg': (XmlToSvgConverter(), tmp_path / 'sample.svg'),
    }
    for converter, output in outputs.values():
        assert converter.convert(str(source), str(output))['success']
        content = output.read_text(encoding='utf-8')
        assert '第一段正文' in content
        assert '第二段正文' in content
        assert '&lt;para&gt;' not in content
        assert '<para>' not in content


def test_json_display_exports_keep_text_and_remove_json_syntax(tmp_path):
    source = tmp_path / 'sample.json'
    source.write_text(
        '{"title":"转换标题","node_type":"paragraph","content":"正文内容"}',
        encoding='utf-8',
    )

    outputs = {
        'html': (JsonToHtmlConverter(), tmp_path / 'sample.html'),
        'svg': (JsonToSvgConverter(), tmp_path / 'sample.svg'),
    }
    for converter, output in outputs.values():
        assert converter.convert(str(source), str(output))['success']
        content = output.read_text(encoding='utf-8')
        assert '转换标题' in content
        assert '正文内容' in content
        assert 'node_type' not in content
        assert '&quot;content&quot;' not in content


def test_json_content_exports_to_yaml_xml_base64_and_xlsx_without_source_code(tmp_path):
    source = tmp_path / 'sample.json'
    source.write_text(
        json.dumps({
            'metadata': {'node_type': 'paragraph', 'tag': 'source-tag'},
            'title': '转换标题',
            'children': [{'content': '正文内容'}, {'content': '第二段'}],
        }, ensure_ascii=False),
        encoding='utf-8',
    )

    yaml_output = tmp_path / 'sample.yaml'
    xml_output = tmp_path / 'sample.xml'
    base64_output = tmp_path / 'sample.base64'
    xlsx_output = tmp_path / 'sample.xlsx'

    assert JsonToYamlConverter().convert(str(source), str(yaml_output))['success']
    assert JsonToXmlConverter().convert(str(source), str(xml_output))['success']
    assert JsonToBase64Converter().convert(str(source), str(base64_output))['success']
    assert JsonToXlsxConverter().convert(str(source), str(xlsx_output))['success']

    yaml_text = yaml_output.read_text(encoding='utf-8')
    xml_text = xml_output.read_text(encoding='utf-8')
    decoded_text = base64.b64decode(base64_output.read_text(encoding='utf-8')).decode('utf-8')
    workbook = load_workbook(xlsx_output, data_only=True)
    try:
        xlsx_values = [row[0] for row in workbook.active.iter_rows(min_row=2, values_only=True)]
    finally:
        workbook.close()

    for content in (yaml_text, xml_text, decoded_text, '\n'.join(map(str, xlsx_values))):
        assert '转换标题' in content
        assert '正文内容' in content
        assert '第二段' in content
        assert 'node_type' not in content
        assert 'metadata' not in content
        assert 'children' not in content
        assert 'source-tag' not in content


def test_html_print_blank_tail_page_is_not_exported(tmp_path):
    source = tmp_path / 'sample.html'
    pdf_output = tmp_path / 'sample.pdf'
    image_output = tmp_path / 'sample.png'
    source.write_text(
        '<!doctype html><meta charset="utf-8"><p>有效正文</p>'
        '<div style="break-before: page; page-break-before: always"></div>',
        encoding='utf-8',
    )

    assert HtmlToPdfConverter().convert(str(source), str(pdf_output))['success']
    result = PdfToImageConverter().convert(str(pdf_output), str(image_output), merge=True)

    assert result['success']
    assert result['page_count'] == 1


def test_html_page_options_are_added_to_browser_print_css():
    html = HtmlToPdfConverter()._prepare_html_content(
        '<html><head></head><body>正文</body></html>',
        {'page_size': 'A3', 'orientation': 'landscape'},
    )

    assert '@page { size: A3 landscape; margin: 10mm; }' in html


def test_text_html_and_browser_print_keep_configured_background(tmp_path):
    source = tmp_path / 'sample.xml'
    output = tmp_path / 'sample.html'
    source.write_text('<document><para>正文</para></document>', encoding='utf-8')
    assert XmlToHtmlConverter().convert(str(source), str(output), background_color='#ddeeff')['success']
    xml_html = output.read_text(encoding='utf-8')
    html = HtmlToPdfConverter()._prepare_html_content(
        '<html><head></head><body>正文</body></html>',
        {'background_color': '#ddeeff'},
    )

    assert '#ddeeff' in xml_html
    assert 'print-color-adjust: exact' in xml_html
    assert '#ddeeff' in html
    assert 'print-color-adjust: exact' in html


def test_excel_html_fit_width_uses_wrapped_full_width_table(tmp_path):
    from openpyxl import Workbook

    source = tmp_path / 'sample.xlsx'
    output = tmp_path / 'sample.html'
    workbook = Workbook()
    workbook.active.append(['第一列', '第二列'])
    workbook.save(source)

    result = ExcelToHtmlConverter().convert(
        str(source),
        str(output),
        excel_scale_mode='fit_width',
    )
    content = output.read_text(encoding='utf-8')

    assert result['success']
    assert 'width: 100%' in content
    assert 'overflow-wrap: anywhere' in content
